"""Turn uploaded or on-disk files into plain text."""
from __future__ import annotations

import io
import json
import re
from pathlib import Path

from .models import Document, Lead

SUPPORTED = {".txt", ".md", ".pdf", ".docx"}


def decode_bytes(data: bytes) -> str:
    """UTF-8 (with or without a Windows BOM) first, then Windows-1252 for files saved by legacy editors."""
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError:
        return data.decode("cp1252", errors="replace")


def normalize_newlines(text: str) -> str:
    """Windows CRLF and old-Mac CR -> LF so paragraph and bullet parsing behave the same on every OS."""
    return text.replace("\r\n", "\n").replace("\r", "\n")


def base_name(name: str) -> str:
    """File name without directories, splitting on both / and \\ (uploads may carry either)."""
    return re.split(r"[\\/]", name)[-1]


def extract_text(name: str, data: bytes) -> str:
    suffix = Path(base_name(name)).suffix.lower()
    if suffix == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(data))
        text = "\n".join(page.extract_text() or "" for page in reader.pages)
    elif suffix == ".docx":
        from docx import Document as DocxDocument

        doc = DocxDocument(io.BytesIO(data))
        text = "\n".join(p.text for p in doc.paragraphs)
    else:
        text = decode_bytes(data)
    return normalize_newlines(text)


def lead_from_files(files: dict[str, bytes]) -> Lead:
    """Build a Lead from {filename: bytes}; metadata.json is optional."""
    metadata: dict = {}
    documents: list[Document] = []
    for name, data in sorted(files.items()):
        base = base_name(name)
        if base.lower() == "metadata.json":
            metadata = json.loads(decode_bytes(data))
        elif Path(base).suffix.lower() in SUPPORTED:
            text = extract_text(base, data).strip()
            if text:
                documents.append(Document(base, text))
    if not documents:
        raise ValueError("Lead packet contains no readable documents.")
    return Lead(metadata, documents)


def load_lead_dir(path: str | Path) -> Lead:
    path = Path(path)
    return lead_from_files({p.name: p.read_bytes() for p in path.iterdir() if p.is_file()})
